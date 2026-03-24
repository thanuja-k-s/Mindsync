from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# Create presentation
prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)

# Define colors
DARK_BLUE = RGBColor(25, 51, 102)
ACCENT_COLOR = RGBColor(0, 102, 204)
TEXT_COLOR = RGBColor(51, 51, 51)
WHITE = RGBColor(255, 255, 255)

def add_title_slide(title, subtitle=""):
    """Add a title slide"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = DARK_BLUE
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(9), Inches(1.5))
    title_frame = title_box.text_frame
    title_frame.word_wrap = True
    p = title_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    
    # Subtitle
    if subtitle:
        subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(9), Inches(2))
        subtitle_frame = subtitle_box.text_frame
        subtitle_frame.word_wrap = True
        p = subtitle_frame.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(28)
        p.font.color.rgb = RGBColor(200, 220, 255)
        p.alignment = PP_ALIGN.CENTER

def add_content_slide(title, bullets):
    """Add a slide with title and bullet points"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    
    # Background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
    title_frame = title_box.text_frame
    p = title_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = DARK_BLUE
    
    # Title underline
    line = slide.shapes.add_shape(1, Inches(0.5), Inches(1.15), Inches(9), Inches(0))
    line.line.color.rgb = ACCENT_COLOR
    line.line.width = Pt(3)
    
    # Content
    content_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.4), Inches(8.6), Inches(5.6))
    text_frame = content_box.text_frame
    text_frame.word_wrap = True
    
    for i, bullet in enumerate(bullets):
        if i > 0:
            text_frame.add_paragraph()
        p = text_frame.paragraphs[i]
        p.text = bullet
        p.font.size = Pt(18)
        p.font.color.rgb = TEXT_COLOR
        p.level = 0
        p.space_before = Pt(8)
        p.space_after = Pt(8)

def add_two_column_slide(title, left_content, right_content):
    """Add a slide with two columns"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
    title_frame = title_box.text_frame
    p = title_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = DARK_BLUE
    
    # Title underline
    line = slide.shapes.add_shape(1, Inches(0.5), Inches(1.15), Inches(9), Inches(0))
    line.line.color.rgb = ACCENT_COLOR
    line.line.width = Pt(3)
    
    # Left column
    left_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(4.3), Inches(5.7))
    left_frame = left_box.text_frame
    left_frame.word_wrap = True
    for i, item in enumerate(left_content):
        if i > 0:
            left_frame.add_paragraph()
        p = left_frame.paragraphs[i]
        p.text = item
        p.font.size = Pt(16)
        p.font.color.rgb = TEXT_COLOR
        p.space_before = Pt(6)
        p.space_after = Pt(6)
    
    # Right column
    right_box = slide.shapes.add_textbox(Inches(5.2), Inches(1.4), Inches(4.3), Inches(5.7))
    right_frame = right_box.text_frame
    right_frame.word_wrap = True
    for i, item in enumerate(right_content):
        if i > 0:
            right_frame.add_paragraph()
        p = right_frame.paragraphs[i]
        p.text = item
        p.font.size = Pt(16)
        p.font.color.rgb = TEXT_COLOR
        p.space_before = Pt(6)
        p.space_after = Pt(6)

# SLIDE 1: Title
add_title_slide(
    "MindSync",
    "AI-Powered Journaling Platform with Local RAG System"
)

# SLIDE 2: Problem Statement
add_two_column_slide(
    "Problem Statement",
    [
        "❌ Before (Traditional AI):",
        "• Slow (2-5 seconds)",
        "• Expensive API costs",
        "• Privacy concerns",
        "• External dependency",
        "• Generic responses"
    ],
    [
        "✅ After (MindSync RAG):",
        "• Fast (100-200ms)",
        "• Free (no API costs)",
        "• Private (local data)",
        "• Fully independent",
        "• Personalized responses"
    ]
)

# SLIDE 3: What is RAG?
add_content_slide(
    "What is RAG?",
    [
        "RAG = Retrieval-Augmented Generation",
        "",
        "🔍 Retrieval: Find relevant journal entries",
        "📚 Augment: Use them as context",
        "✍️ Generation: Create personalized response",
        "",
        "Example: User asks 'How am I doing?'",
        "→ AI finds 5 related entries",
        "→ Responds with context: 'You're showing enthusiasm about your project!'"
    ]
)

# SLIDE 4: Core Features
add_content_slide(
    "Core Features",
    [
        "✅ User Authentication - Secure JWT & bcryptjs",
        "✅ Journal Writing - Create/edit entries with mood & tags",
        "✅ MemoTalks AI - Local RAG-powered companion",
        "✅ Daily Streak - Track consecutive journaling days",
        "✅ Goals Tracking - Health, Career, Finance, Personal",
        "✅ Reminders - Priority-based notifications",
        "✅ Insights Dashboard - Mood trends & analytics",
        "✅ Dark/Light Theme - User preferences",
        "✅ Media Support - Upload images & files"
    ]
)

# SLIDE 5: Technical Architecture
add_content_slide(
    "Technical Architecture",
    [
        "React UI (Frontend)",
        "       ↓",
        "Node.js + Express Server",
        "  ├─ RAG Routes (/api/rag/query)",
        "  ├─ Entry Routes (/api/entries)",
        "  └─ RAG Service (embedding & similarity)",
        "       ↓",
        "MongoDB Database",
        "  └─ RAGIndex Collection (embeddings)"
    ]
)

# SLIDE 6: Technology Stack
add_two_column_slide(
    "Technology Stack",
    [
        "Frontend:",
        "• React 18.2.0",
        "• React Router DOM 7.9.6",
        "• Chart.js & react-chartjs-2",
        "• Crypto-JS (encryption)",
        "• CSS3 (responsive)",
    ],
    [
        "Backend:",
        "• Node.js + Express 4.18.2",
        "• MongoDB 7.0",
        "• Mongoose ORM",
        "• JWT authentication",
        "• bcryptjs (password hashing)"
    ]
)

# SLIDE 7: AI/ML Implementation
add_content_slide(
    "AI/ML Implementation",
    [
        "🤖 RAG System Features:",
        "",
        "📊 TF-IDF Text Embedding",
        "  → Converts text to 384-dimensional vectors",
        "",
        "🎯 Cosine Similarity Matching",
        "  → Finds most relevant journal entries",
        "",
        "⚡ Real-time Indexing",
        "  → Automatically indexes new entries",
        "",
        "💭 Emotion Analysis",
        "  → Understands context and mood"
    ]
)

# SLIDE 8: Database Schema
add_content_slide(
    "Database Schema",
    [
        "📌 Users Collection",
        "  → Store usernames, emails, hashed passwords, streaks",
        "",
        "📔 Entries Collection",
        "  → Journal entries with mood, tags, timestamps, media",
        "",
        "🎯 Goals Collection",
        "  → Personal goals with categories and progress",
        "",
        "⏰ Reminders Collection",
        "  → Reminders with due dates and priorities",
        "",
        "🧠 RAGIndex Collection",
        "  → Entry embeddings (384-dim vectors) + metadata"
    ]
)

# SLIDE 9: API Endpoints
add_content_slide(
    "API Endpoints",
    [
        "🔐 Authentication:",
        "  POST /api/auth/login  |  POST /api/auth/signup",
        "",
        "📝 Entry Management:",
        "  POST/GET/PUT/DELETE /api/entries",
        "",
        "🎯 Goals:",
        "  POST/GET/PUT/DELETE /api/goals",
        "",
        "⏰ Reminders:",
        "  POST/GET/PUT/DELETE /api/reminders",
        "",
        "🧠 MemoTalks AI:",
        "  POST /api/rag/query    (personalized responses)"
    ]
)

# SLIDE 10: Application Pages
add_content_slide(
    "Application Pages & Routes",
    [
        "🏠 Landing (/) - Welcome page with feature overview",
        "🔐 Login (/auth) - User authentication",
        "📝 Journaling (/journal) - Write and edit entries",
        "📚 Entries (/entries) - View all entries, search, filter",
        "🧠 MemoTalks (/sage) - AI companion chat",
        "📊 Insights (/insights) - Mood trends & analytics",
        "🎯 Goals (/goals) - Goal management",
        "⏰ Reminders (/reminders) - Reminder management",
        "⚙️ Settings (/settings) - Theme & preferences"
    ]
)

# SLIDE 11: Data Flow
add_content_slide(
    "RAG System Data Flow",
    [
        "1️⃣ User writes journal entry",
        "",
        "2️⃣ Entry automatically indexed",
        "   → Convert to TF-IDF vector (384-dim)",
        "   → Store embedding in RAGIndex",
        "",
        "3️⃣ User asks question (MemoTalks)",
        "",
        "4️⃣ Find similar entries",
        "   → Convert question to vector",
        "   → Calculate cosine similarity",
        "   → Return top 5 matches",
        "",
        "5️⃣ Generate response using context",
        "   → Analyze emotions & patterns",
        "   → Create personalized answer"
    ]
)

# SLIDE 12: Security Features
add_content_slide(
    "Security & Privacy",
    [
        "🔒 JWT Token-based Authentication",
        "  → Secure session management",
        "",
        "🔐 bcryptjs Password Hashing",
        "  → Passwords never stored in plaintext",
        "",
        "👤 User Data Isolation",
        "  → Users can only access their own data",
        "",
        "🛡️ CORS Protection",
        "  → Cross-origin requests validated",
        "",
        "✅ Input Validation",
        "  → All inputs sanitized and validated",
        "",
        "🔑 Environment Variables",
        "  → Sensitive data secured in .env files"
    ]
)

# SLIDE 13: Project Structure
add_content_slide(
    "Project Structure",
    [
        "MindSync/",
        "├── src/ (React Frontend)",
        "│   ├── components/ (Header, Nav, Sidebar, Footer)",
        "│   ├── pages/ (Auth, Journal, Entries, Goals, etc.)",
        "│   ├── contexts/ (State management)",
        "│   └── styles/ (Theme, CSS)",
        "├── server/ (Node.js Backend)",
        "│   ├── models/ (User, Entry, Goal, RAGIndex)",
        "│   ├── routes/ (auth, entries, goals, rag)",
        "│   └── utils/ (embeddingService, ragService)",
        "└── build/ (Production optimized frontend)"
    ]
)

# SLIDE 14: Key Achievements
add_content_slide(
    "Key Achievements",
    [
        "✅ Complete RAG AI system operational",
        "✅ Zero external API dependencies",
        "✅ Real-time analytics dashboard",
        "✅ Secure authentication system (JWT + bcryptjs)",
        "✅ Responsive UI with dark/light modes",
        "✅ Full CRUD operations for all features",
        "✅ Media upload support (images & files)",
        "✅ Daily streak tracking system",
        "✅ Emotion-aware AI responses"
    ]
)

# SLIDE 15: Technology Highlights
add_content_slide(
    "Technology Highlights",
    [
        "🚀 MERN Stack (MongoDB, Express, React, Node.js)",
        "🤖 Local AI with TF-IDF & Cosine Similarity",
        "📊 Real-time Data Visualization (Chart.js)",
        "🔐 Advanced Security (JWT, bcryptjs, CORS)",
        "🎨 Responsive Design (Mobile-friendly)",
        "⚡ Fast Performance (100-200ms AI responses)",
        "🔄 RESTful API Architecture",
        "💾 MongoDB for scalable data storage"
    ]
)

# SLIDE 16: Future Enhancements
add_content_slide(
    "Future Enhancements",
    [
        "🎙️ Voice-to-journal integration",
        "📥 Export entries as PDF/CSV",
        "📱 Native mobile app (React Native/Flutter)",
        "🌐 Social sharing features",
        "🔮 Advanced predictive analytics",
        "🌍 Multi-language support",
        "🎯 Personalized recommendations",
        "📈 Advanced ML models (sentiment analysis)"
    ]
)

# SLIDE 17: Deployment & Ports
add_content_slide(
    "Deployment Information",
    [
        "🖥️ Frontend Server",
        "  → Port: 3000/3001",
        "  → Technology: React (npm start)",
        "",
        "🔧 Backend Server",
        "  → Port: 3002",
        "  → Technology: Node.js + Express",
        "",
        "💾 Database",
        "  → MongoDB 7.0 (local or Atlas cloud)",
        "",
        "📦 Build Command",
        "  → npm run build (optimized React build)"
    ]
)

# SLIDE 18: Conclusion
slide = prs.slides.add_slide(prs.slide_layouts[6])
background = slide.background
fill = background.fill
fill.solid()
fill.fore_color.rgb = DARK_BLUE

# Main text
conclusion_box = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(9), Inches(3.5))
conclusion_frame = conclusion_box.text_frame
conclusion_frame.word_wrap = True

p = conclusion_frame.paragraphs[0]
p.text = "MindSync"
p.font.size = Pt(48)
p.font.bold = True
p.font.color.rgb = WHITE
p.alignment = PP_ALIGN.CENTER

conclusion_frame.add_paragraph()
p = conclusion_frame.paragraphs[1]
p.text = "Intelligent, Private, and Personal Journaling"
p.font.size = Pt(24)
p.font.color.rgb = RGBColor(200, 220, 255)
p.alignment = PP_ALIGN.CENTER
p.space_before = Pt(12)

conclusion_frame.add_paragraph()
p = conclusion_frame.paragraphs[2]
p.text = "Local RAG AI • Secure Authentication • Real-time Analytics"
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(180, 200, 255)
p.alignment = PP_ALIGN.CENTER
p.space_before = Pt(20)

# Save presentation
prs.save('ppt/MindSync_Complete_Presentation.pptx')
print("✅ PowerPoint presentation created successfully!")
print("📁 File: ppt/MindSync_Complete_Presentation.pptx")
print("📊 Total slides: 18")
